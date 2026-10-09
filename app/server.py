import asyncio
from contextlib import asynccontextmanager
from datetime import date
import hashlib
from io import BytesIO
import os
from pathlib import Path
import secrets
from urllib.parse import urlsplit
import uuid

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field

from .service import Service

STATIC = Path(__file__).parent / "static"
ACTIONS = {
    "grid_layout",
    "market_watch",
    "preview",
    "create",
    "edit",
    "archive",
    "start",
    "stop",
    "quote",
    "reconcile",
    "login",
    "select_account",
    "lookup",
    "logout",
    "fee",
    "opening_cost",
    "backup",
    "report",
    "import_legacy",
    "shutdown",
    "live_control",
    "reconnect",
    "order_diagnostics",
    "order_evidence",
    "retry_order",
    "recovery_preview",
    "recovery_confirm",
    "bulk_stop",
    "ui_preferences",
}


def runtime_path():
    base = Path(os.environ.get("LOCALAPPDATA", str(Path.home())))
    return Path(os.environ.get("KGI_GRID_RUNTIME", base / "KGI_US_Grid_Trading_Web" / "runtime"))


def runtime_identity(directory):
    normalized = os.path.normcase(os.path.abspath(directory))
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


class RuntimeLock:
    def __init__(self, path):
        self.path = path
        self.handle = None

    def acquire(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = open(self.path, "a+b")
        self.handle.seek(0, 2)
        if self.handle.tell() == 0:
            self.handle.write(b"0")
            self.handle.flush()
        self.handle.seek(0)
        try:
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            self.handle.close()
            self.handle = None
            raise RuntimeError("這個資料目錄已有工作台使用中") from None

    def release(self):
        if self.handle:
            self.handle.close()
            self.handle = None


class Command(BaseModel):
    model_config = ConfigDict(extra="forbid")
    request_id: str = Field(min_length=1, max_length=80)
    payload: dict = Field(default_factory=dict)


def create_app(path=None, service_factory=Service):
    directory = Path(path) if path else runtime_path()
    service = service_factory(directory / "trading.sqlite3")
    guard = RuntimeLock(directory / "instance.lock")
    token = secrets.token_urlsafe(32)
    instance = uuid.uuid4().hex

    @asynccontextmanager
    async def lifespan(app):
        guard.acquire()
        try:
            service.start()
            yield
        finally:
            try:
                service.close()
            finally:
                # A blocked actor must retain custody of its DB lock until process exit.
                if not service.thread.is_alive() and service.cleanup_complete.is_set():
                    guard.release()

    app = FastAPI(
        title="KGI 美股網格工作台", lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None
    )
    app.state.service = service
    app.state.instance = instance
    app.state.shutdown_callback = None

    @app.middleware("http")
    async def local_boundary(request: Request, call_next):
        host = request.headers.get("host", "")
        try:
            parsed = urlsplit("http://" + host)
            valid = parsed.hostname in ("127.0.0.1", "localhost") and parsed.username is None
        except ValueError:
            valid = False
        origin = request.headers.get("origin")
        if (
            not valid
            or (origin and origin != f"http://{host}")
            or request.headers.get("sec-fetch-site") == "cross-site"
        ):
            return JSONResponse({"detail": "僅允許本機同來源操作"}, status_code=403)
        if request.method not in ("GET", "HEAD"):
            if not secrets.compare_digest(request.headers.get("x-grid-token", ""), token):
                return JSONResponse({"detail": "操作權杖失效，請重新整理"}, status_code=403)
            body = await request.body()
            if len(body) > 1_000_000:
                return JSONResponse({"detail": "請求過大"}, status_code=413)
        response = await call_next(request)
        response.headers["Cache-Control"] = "no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"
        )
        return response

    @app.get("/")
    def home():
        return FileResponse(STATIC / "index.html")

    @app.get("/api/session")
    def session():
        return {"token": token, "instance": instance, "version": "0.1.4"}

    @app.get("/api/health")
    def health():
        # Startup identity only; no credentials, account data or actor commands.
        return {
            "product": "kgi-us-grid-web",
            "protocol": 1,
            "instance": instance,
            "version": "0.1.4",
            "pid": os.getpid(),
            "runtime_id": runtime_identity(directory),
            "ready": service.thread.is_alive() and not service.stop_event.is_set(),
        }

    @app.get("/api/state")
    def state():
        return service.get_state()

    @app.get("/api/diagnostics")
    def diagnostics():
        result = service.audit.snapshot()
        result["actor_alive"] = service.thread.is_alive()
        return result

    @app.get("/api/audit/download")
    def audit_download():
        # Only the application's redacted audit; never vendor logs, credentials or DB.
        with service.audit.lock:
            data = service.audit.path.read_bytes()
        return Response(
            data,
            media_type="application/x-ndjson",
            headers={"Content-Disposition": f'attachment; filename="{service.audit.path.name}"'},
        )

    @app.post("/api/commands/{action}", status_code=202)
    def command(action: str, body: Command):
        if action not in ACTIONS:
            raise HTTPException(404, "未知操作")
        try:
            cid = service.submit(action, body.payload, body.request_id)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from None
        return {"command_id": cid, "status": "queued"}

    @app.get("/api/commands/{cid}")
    async def command_result(cid: str):
        result = service.get_result(cid)
        if not result:
            raise HTTPException(404, "操作紀錄已過期；請查詢帳本，不要重送交易")
        if result.get("result", {}).get("shutdown") and app.state.shutdown_callback:
            asyncio.get_running_loop().call_later(0.5, app.state.shutdown_callback)
        return result

    @app.get("/api/export.xlsx")
    async def export(day: str | None = None):
        if day:
            try:
                date.fromisoformat(day)
            except ValueError:
                raise HTTPException(422, "日期格式需為 YYYY-MM-DD") from None
        cid = service.submit("report", {"day": day}, uuid.uuid4().hex)
        for _ in range(100):
            result = service.get_result(cid)
            if result["status"] in ("done", "error"):
                break
            await asyncio.sleep(0.1)
        if result["status"] != "done":
            raise HTTPException(503, "報表尚未可用；請稍後再查")
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        report = result["result"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "損益與目前持倉"
        sheet.append(
            ["日期範圍", day or "全部", "成交日期時區", "America/New_York", "成本法", "移動加權平均"]
        )
        sheet.append(
            ["注意", "持倉與成本為目前值；成交時間不完整時成本依接收順序暫估、單日損益未知。非券商對帳單。"]
        )
        sheet.append(
            [
                "策略",
                "模式",
                "商品",
                "成本幣別",
                "目前股數",
                "目前總成本",
                "目前均價",
                "區間已實現",
                "費用完整",
                "版本",
                "成交時間完整",
                "成本順序暫估",
                "目前設定結算方式",
                "歷史成交結算方式",
            ]
        )
        for row in report["rows"]:
            values = [
                row[k]
                for k in (
                    "name",
                    "mode",
                    "symbol",
                    "currency",
                    "quantity",
                    "cost",
                    "average_cost",
                    "realized",
                    "fees_complete",
                )
            ] + [
                report["version"],
                row["time_complete"],
                row["cost_provisional"],
                row["current_settlement_currency"],
                ", ".join(row["historical_settlement_currencies"]),
            ]
            sheet.append(
                [
                    ("'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v)
                    for v in values
                ]
            )
        for cell in sheet[3]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="163D5A")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
        sheet.freeze_panes = "A4"
        for title, headers, keys, records in (
            (
                "結算方式彙總",
                [
                    "策略",
                    "商品",
                    "送單結算方式",
                    "買入股數",
                    "賣出股數",
                    "買入金額 USD",
                    "賣出金額 USD",
                    "來源日期完整",
                ],
                [
                    "name",
                    "symbol",
                    "settlement_currency",
                    "buy_quantity",
                    "sell_quantity",
                    "buy_value_usd",
                    "sell_value_usd",
                    "time_complete",
                ],
                report["settlement_rows"],
            ),
            (
                "成交明細",
                [
                    "策略",
                    "商品",
                    "委託識別",
                    "送單結算方式",
                    "方向",
                    "股數",
                    "成交金額 USD",
                    "記錄時間",
                    "時間品質",
                    "來源時間",
                ],
                [
                    "name",
                    "symbol",
                    "order_id",
                    "settlement_currency",
                    "side",
                    "quantity",
                    "value",
                    "time",
                    "time_quality",
                    "source_time",
                ],
                report["fill_details"],
            ),
        ):
            extra = workbook.create_sheet(title)
            extra.append(
                [
                    "說明",
                    "TWD=台幣結算；MUT=外幣結算。金額皆為 USD，未換匯或推估實際扣款。單日報表排除來源日期未知成交。",
                ]
            )
            extra.append(headers)
            for record in records:
                values = [record.get(k, "") for k in keys]
                extra.append(
                    [
                        "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v
                        for v in values
                    ]
                )
            extra.freeze_panes = "A3"
            for cell in extra[2]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="163D5A")
            for col in extra.columns:
                extra.column_dimensions[col[0].column_letter].width = 24
        output = BytesIO()
        workbook.save(output)
        return Response(
            output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="grid-pnl-{day or "all"}.xlsx"'},
        )

    app.mount("/static", StaticFiles(directory=STATIC), name="static")
    return app
